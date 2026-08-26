"""Drawing primitives: the only module that touches python-pptx shape APIs.

Every primitive takes inches and token-resolved colours, names the shape it creates so
the verifier can reason about it, and defends against the traps that silently produce
an off-brand deck — inherited theme shadows, schema-ordered table borders, chevron
notches eating labels, and masters that never clone their slide-number placeholder.
"""

from __future__ import annotations

import uuid
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, replace

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

from .tokens import Theme

__all__ = [
    "DF",
    "Box",
    "TextBlock",
    "add_bullets",
    "add_connector",
    "add_gradient_bar",
    "add_shape",
    "add_slide_number",
    "add_table",
    "add_textbox",
    "name_shape",
    "set_notes",
    "shape_role",
    "write_blocks",
]

DF = "df:"
"""Prefix marking a shape this package drew, as opposed to an inherited placeholder."""

_ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_ALIGN_XML = {"left": "l", "center": "ctr", "right": "r", "justify": "just"}
# "No Style, No Grid" — without it PowerPoint's default blue table style fights
# explicit cell fills.
_PLAIN_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def name_shape(role: str, detail: str = "") -> str:
    """Build a shape name the verifier can parse: ``df:role`` or ``df:role:detail``."""
    return f"{DF}{role}" + (f":{detail}" if detail else "")


def shape_role(name: str) -> str:
    """Extract the role from a generated shape name; empty for foreign shapes."""
    text = str(name)
    if not text.startswith(DF):
        return ""
    return text[len(DF) :].split(":")[0]


@dataclass(frozen=True)
class Box:
    """A rectangle in inches. All layout maths happens here, never in EMU."""

    left: float
    top: float
    width: float
    height: float

    @property
    def right(self) -> float:
        """Right edge."""
        return self.left + self.width

    @property
    def bottom(self) -> float:
        """Bottom edge."""
        return self.top + self.height

    def inset(self, pad: float, pad_y: float | None = None) -> Box:
        """Shrink on all sides by ``pad`` (and ``pad_y`` vertically when given)."""
        vertical = pad if pad_y is None else pad_y
        return Box(
            self.left + pad, self.top + vertical, self.width - 2 * pad, self.height - 2 * vertical
        )

    def columns(self, count: int, gap: float) -> list[Box]:
        """Split into ``count`` equal columns separated by ``gap``."""
        width = (self.width - gap * (count - 1)) / count
        return [
            Box(self.left + i * (width + gap), self.top, width, self.height) for i in range(count)
        ]

    def rows(self, count: int, gap: float) -> list[Box]:
        """Split into ``count`` equal rows separated by ``gap``."""
        height = (self.height - gap * (count - 1)) / count
        return [
            Box(self.left, self.top + i * (height + gap), self.width, height) for i in range(count)
        ]

    def top_slice(self, height: float) -> Box:
        """The top ``height`` inches of this box."""
        return replace(self, height=height)

    def bottom_slice(self, height: float) -> Box:
        """The bottom ``height`` inches of this box."""
        return Box(self.left, self.bottom - height, self.width, height)

    def shrink_bottom(self, amount: float) -> Box:
        """Same box with ``amount`` inches removed from the bottom."""
        return replace(self, height=self.height - amount)

    def moved(self, *, dx: float = 0.0, dy: float = 0.0) -> Box:
        """Same box translated by ``dx`` / ``dy``."""
        return replace(self, left=self.left + dx, top=self.top + dy)

    def as_emu(self) -> tuple[Emu, Emu, Emu, Emu]:
        """Left, top, width, height as EMU, ready for python-pptx."""
        return (Inches(self.left), Inches(self.top), Inches(self.width), Inches(self.height))


@dataclass(frozen=True)
class TextBlock:
    """One paragraph: what it says and which rung of the type scale it sits on."""

    text: str
    step: str = "body"
    color: str = "body_text"
    align: str = "left"
    bullet: str = ""
    bullet_color: str = ""
    indent_level: int = 0
    bold: bool | None = None
    space_after_pt: float | None = None
    line: float | None = None


def rgb(hex_value: str) -> RGBColor:
    """Convert a token hex string to a python-pptx colour."""
    return RGBColor.from_string(hex_value.lstrip("#").upper())


def kill_shadow(shape) -> None:
    """Remove the inherited theme effects that make generated diagrams look like clip art.

    Two things have to go. ``shadow.inherit = False`` writes an empty ``a:effectLst``,
    which PowerPoint honours. The ``p:style`` element python-pptx attaches to every new
    autoshape also points at the theme's effect, fill and line references, and some
    renderers apply that in preference — so it is dropped outright. Fill and outline are
    always set explicitly here, so nothing is lost.
    """
    try:
        shape.shadow.inherit = False
    except (AttributeError, NotImplementedError):
        pass
    element = getattr(shape, "_element", None)
    if element is None:
        return
    style = element.find(qn("p:style"))
    if style is not None:
        element.remove(style)


def _prepare_frame(shape, *, anchor: str, wrap: bool, margin: float = 0.0) -> None:
    frame = shape.text_frame
    frame.word_wrap = wrap
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = _ANCHOR[anchor]
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, Inches(margin))


def _paragraph_props(
    paragraph,
    theme: Theme,
    block: TextBlock,
    *,
    size_pt: float,
    line: float,
    space_after_pt: float,
) -> None:
    """Replace ``a:pPr`` wholesale so every child lands in schema order."""
    marl = 0
    indent = 0
    if block.bullet:
        marl = Inches(
            theme.spacing.bullet_indent_in + block.indent_level * theme.spacing.sub_indent_in
        )
        indent = -Inches(theme.spacing.bullet_indent_in)
    elif block.indent_level:
        marl = Inches(block.indent_level * theme.spacing.sub_indent_in)
    bullet_xml = (
        f'<a:buClr><a:srgbClr val="{theme.color(block.bullet_color or block.color)}"/></a:buClr>'
        f'<a:buSzPct val="90000"/>'
        f'<a:buFont typeface="Arial"/>'
        f'<a:buChar char="{block.bullet}"/>'
        if block.bullet
        else "<a:buNone/>"
    )
    xml = (
        f'<a:pPr {nsdecls("a")} marL="{int(marl)}" indent="{int(indent)}" '
        f'algn="{_ALIGN_XML[block.align]}">'
        f'<a:lnSpc><a:spcPct val="{int(line * 100000)}"/></a:lnSpc>'
        f'<a:spcBef><a:spcPts val="0"/></a:spcBef>'
        f'<a:spcAft><a:spcPts val="{int(space_after_pt * 100)}"/></a:spcAft>'
        f"{bullet_xml}"
        "</a:pPr>"
    )
    existing = paragraph._p.find(qn("a:pPr"))
    if existing is not None:
        paragraph._p.remove(existing)
    paragraph._p.insert(0, parse_xml(xml))
    paragraph.alignment = _ALIGN[block.align]


def write_blocks(shape, theme: Theme, blocks: Sequence[TextBlock]) -> None:
    """Write paragraphs into a shape's text frame, styling every run explicitly.

    Runs never rely on theme inheritance: a drawn text box falls back to the minor font
    at the wrong weight, which is invisible until someone opens the deck on a machine
    without the corporate font installed.
    """
    frame = shape.text_frame
    frame.clear()
    for index, block in enumerate(blocks):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        step = theme.typography.step(block.step)
        size_pt = step.size_pt
        line = block.line if block.line is not None else step.line
        space_after = (
            block.space_after_pt if block.space_after_pt is not None else step.space_after_pt
        )
        _paragraph_props(
            paragraph, theme, block, size_pt=size_pt, line=line, space_after_pt=space_after
        )
        run = paragraph.add_run()
        run.text = block.text.upper() if step.caps else block.text
        run.font.size = Pt(size_pt)
        run.font.bold = step.bold if block.bold is None else block.bold
        run.font.name = theme.typography.fonts[step.font]
        run.font.color.rgb = rgb(theme.color(block.color))
        # Corporate layouts often force all-caps on the title placeholder; case is a
        # content decision, so it is neutralised and applied here instead.
        run.font._rPr.set("cap", "none")


def add_textbox(
    slide,
    box: Box,
    theme: Theme,
    blocks: Sequence[TextBlock],
    *,
    role: str = "text",
    detail: str = "",
    anchor: str = "top",
    wrap: bool = True,
):
    """Add a plain text box carrying pre-styled paragraphs."""
    shape = slide.shapes.add_textbox(*box.as_emu())
    shape.name = name_shape(role, detail)
    _prepare_frame(shape, anchor=anchor, wrap=wrap)
    write_blocks(shape, theme, blocks)
    kill_shadow(shape)
    return shape


def add_bullets(
    slide,
    box: Box,
    theme: Theme,
    items: Iterable[str],
    *,
    step: str = "body",
    color: str = "body_text",
    bullet: str = "\u2013",
    bullet_color: str = "",
    anchor: str = "top",
    role: str = "bullets",
):
    """Add a bullet block with a hanging indent.

    ``bullet_color`` carries the slide's semantic entity colour. The type itself stays
    the theme's body colour: category is signalled by the glyph, never by hue-coloured
    text.
    """
    blocks = [
        TextBlock(text=item, step=step, color=color, bullet=bullet, bullet_color=bullet_color)
        for item in items
    ]
    return add_textbox(slide, box, theme, blocks, role=role, anchor=anchor)


def add_shape(
    slide,
    shape_type: MSO_SHAPE,
    box: Box,
    *,
    role: str,
    detail: str = "",
    fill: str | None = None,
    line: str | None = None,
    line_pt: float = 1.0,
    adjustment: float | None = None,
):
    """Add an autoshape with fill, outline and adjustment set explicitly."""
    shape = slide.shapes.add_shape(shape_type, *box.as_emu())
    shape.name = name_shape(role, detail)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_pt)
    if adjustment is not None and shape.adjustments:
        shape.adjustments[0] = adjustment
    shape.text_frame.text = ""
    _prepare_frame(shape, anchor="middle", wrap=True, margin=0.04)
    kill_shadow(shape)
    return shape


def add_gradient_bar(slide, box: Box, theme: Theme, start: str, end: str, *, role: str = "accent"):
    """Add the brand accent bar: a thin rectangle with a horizontal gradient."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *box.as_emu())
    shape.name = name_shape(role)
    shape.line.fill.background()
    fill = shape.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = rgb(theme.color(start))
    stops[0].position = 0.0
    stops[-1].color.rgb = rgb(theme.color(end))
    stops[-1].position = 1.0
    fill.gradient_angle = 0.0
    kill_shadow(shape)
    return shape


def add_connector(
    slide,
    start: tuple[float, float],
    end: tuple[float, float],
    *,
    color: str,
    width_pt: float = 1.5,
):
    """Add a straight connector between two points in inches."""
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(start[0]),
        Inches(start[1]),
        Inches(end[0]),
        Inches(end[1]),
    )
    connector.name = name_shape("connector")
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width_pt)
    kill_shadow(connector)
    return connector


def add_slide_number(slide, theme: Theme, *, color: str = "secondary_text") -> None:
    """Stamp a live ``slidenum`` field on the slide.

    python-pptx does not clone the master's slide-number placeholder, so a generated
    deck has no page numbers at all unless the field is injected by hand.
    """
    grid = theme.grid
    box = Box(
        left=theme.canvas.width_in - grid.margin_right_in - 1.0,
        top=grid.footer_top_in,
        width=1.0,
        height=grid.footer_height_in,
    )
    shape = slide.shapes.add_textbox(*box.as_emu())
    shape.name = name_shape("pagenum")
    _prepare_frame(shape, anchor="middle", wrap=False)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    step = theme.typography.step("tiny")
    field_xml = (
        f'<a:fld {nsdecls("a")} id="{{{uuid.uuid4()}}}" type="slidenum">'
        f'<a:rPr lang="en-US" sz="{int(step.size_pt * 100)}" b="1">'
        f'<a:solidFill><a:srgbClr val="{theme.color(color)}"/></a:solidFill>'
        f'<a:latin typeface="{theme.typography.fonts[step.font]}"/></a:rPr>'
        "<a:t>1</a:t></a:fld>"
    )
    paragraph._p.append(parse_xml(field_xml))
    kill_shadow(shape)


def set_notes(slide, text: str) -> None:
    """Write speaker notes, creating the notes slide on demand."""
    slide.notes_slide.notes_text_frame.text = text


def _border_xml(tag: str, color: str | None, width_pt: float) -> str:
    if color is None:
        return f'<a:{tag} {nsdecls("a")}><a:noFill/></a:{tag}>'
    return (
        f'<a:{tag} {nsdecls("a")} w="{int(width_pt * 12700)}" cap="flat" cmpd="sng" algn="ctr">'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:prstDash val="solid"/></a:{tag}>'
    )


def set_cell_borders(
    cell,
    *,
    left: str | None = None,
    right: str | None = None,
    top: str | None = None,
    bottom: str | None = None,
    width_pt: float = 1.0,
) -> None:
    """Write cell borders into ``tcPr``.

    The schema fixes the child order as ``lnL, lnR, lnT, lnB``, so each element is
    inserted at index 0 in reverse; appending instead produces a file PowerPoint
    silently repairs by dropping the borders.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):
        existing = tc_pr.find(qn(tag))
        if existing is not None:
            tc_pr.remove(existing)
    for tag, color in (("lnB", bottom), ("lnT", top), ("lnR", right), ("lnL", left)):
        tc_pr.insert(0, parse_xml(_border_xml(tag, color, width_pt)))


def add_table(
    slide,
    box: Box,
    theme: Theme,
    *,
    rows: int,
    cols: int,
    col_widths: Sequence[float] | None = None,
    detail: str = "",
):
    """Add a native table with the built-in style flags disabled.

    Returns:
        The ``(graphic_frame, table)`` pair; callers style cells through
        :func:`style_cell` and :func:`set_cell_borders`.
    """
    frame = slide.shapes.add_table(rows, cols, *box.as_emu())
    frame.name = name_shape("table", detail)
    table = frame.table
    for flag in ("first_row", "last_row", "first_col", "last_col", "horz_banding", "vert_banding"):
        setattr(table, flag, False)
    style_id = table._tbl.tblPr.find(qn("a:tableStyleId"))
    if style_id is None:
        table._tbl.tblPr.append(
            parse_xml(f'<a:tableStyleId {nsdecls("a")}>{_PLAIN_TABLE_STYLE}</a:tableStyleId>')
        )
    else:
        style_id.text = _PLAIN_TABLE_STYLE
    if col_widths:
        for column, width in zip(table.columns, col_widths):
            column.width = Inches(width)
    row_height = max(box.height / rows, theme.spacing.table_row_min_in)
    for row in table.rows:
        row.height = Inches(row_height)
    return frame, table


def style_cell(
    cell,
    theme: Theme,
    text: str,
    *,
    step: str = "micro",
    color: str = "body_text",
    fill: str | None = None,
    align: str = "left",
    bold: bool | None = None,
) -> None:
    """Fill and type one table cell."""
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(theme.spacing.table_margin_h_in)
    cell.margin_right = Inches(theme.spacing.table_margin_h_in)
    cell.margin_top = Inches(theme.spacing.table_margin_v_in)
    cell.margin_bottom = Inches(theme.spacing.table_margin_v_in)
    write_blocks(
        cell, theme, [TextBlock(text=text, step=step, color=color, align=align, bold=bold)]
    )
