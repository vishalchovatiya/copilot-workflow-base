"""Structural verification that runs everywhere, plus an optional visual pass.

The XML tells you nothing about whether text overflowed its card or a title clipped the
logo, so verification has two halves. The structural half is pure Python and always
runs: bounds, budgets, notes, grid alignment, editability, colour discipline. The
visual half rasterises the deck through :mod:`deckforge.render` and reports which tier
did it, so a caller always knows how much confidence to place in a clean report.
"""

from __future__ import annotations

import json
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from . import measure
from .primitives import DF, shape_role
from .render import RenderResult, render
from .tokens import Theme

__all__ = ["Finding", "VerifyReport", "verify"]

_TOLERANCE_IN = 0.02
_GRID_DRIFT_IN = 0.5
_OVERLAP_OK = frozenset({"chevron", "label", "accent", "artifact", "band", "backdrop"})
_MIN_OVERLAP_IN2 = 0.06
_OUTSIDE_BODY = frozenset({"title", "accent", "kicker", "takeaway", "unknowns", "backdrop"})
_BALANCE_SLACK_IN = 1.0
_RASTER_TYPES = frozenset(
    {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE, MSO_SHAPE_TYPE.MEDIA}
)
_UNKNOWN_TOKENS = ("TBC", "TBD", "???")


@dataclass(frozen=True)
class Finding:
    """One verification result, precise enough to act on without opening the deck."""

    severity: str  # error | warn | info
    check: str
    slide: int | None
    message: str

    def __str__(self) -> str:
        where = f"slide {self.slide}" if self.slide else "deck"
        return f"[{self.severity:<5}] {where:<9} {self.check}: {self.message}"


@dataclass
class VerifyReport:
    """Everything verification learned about a deck."""

    deck: Path
    findings: list[Finding] = field(default_factory=list)
    render_result: RenderResult | None = None
    slide_count: int = 0

    @property
    def errors(self) -> list[Finding]:
        """Findings that must be fixed before shipping."""
        return [f for f in self.findings if f.severity == "error"]

    @property
    def warnings(self) -> list[Finding]:
        """Findings worth a look but not blocking."""
        return [f for f in self.findings if f.severity == "warn"]

    @property
    def ok(self) -> bool:
        """True when nothing blocking was found."""
        return not self.errors

    def to_text(self) -> str:
        """Human-readable report ending in a one-line verdict."""
        tier = self.render_result.tier if self.render_result else "not run"
        lines = [
            f"deck      : {self.deck}",
            f"slides    : {self.slide_count}",
            f"render    : tier={tier}"
            + (
                f" images={len(self.render_result.images)}"
                if self.render_result and self.render_result.images
                else ""
            ),
        ]
        if self.render_result and self.render_result.message:
            lines.append(f"render note: {self.render_result.message}")
        lines.append("")
        lines.extend(str(f) for f in self.findings) or lines.append("no findings")
        lines.append("")
        lines.append(
            f"VERDICT: {'PASS' if self.ok else 'FAIL'} "
            f"({len(self.errors)} errors, {len(self.warnings)} warnings)"
        )
        return "\n".join(lines)

    def to_json(self) -> str:
        """Machine-readable report."""
        return json.dumps(
            {
                "deck": str(self.deck),
                "slides": self.slide_count,
                "render_tier": self.render_result.tier if self.render_result else None,
                "ok": self.ok,
                "findings": [
                    {
                        "severity": f.severity,
                        "check": f.check,
                        "slide": f.slide,
                        "message": f.message,
                    }
                    for f in self.findings
                ],
            },
            indent=2,
        )


def _inches(value) -> float:
    return 0.0 if value is None else Emu(int(value)).inches


def _box(shape) -> tuple[float, float, float, float]:
    left = _inches(shape.left)
    top = _inches(shape.top)
    return left, top, left + _inches(shape.width), top + _inches(shape.height)


def _generated(shapes: Iterable) -> list:
    return [s for s in shapes if str(s.name).startswith(DF)]


def _all_text(shape) -> str:
    if not shape.has_text_frame:
        return ""
    return shape.text_frame.text


def _check_bounds(index: int, slide, theme: Theme) -> Iterator[Finding]:
    width, height = theme.canvas.width_in, theme.canvas.height_in
    for shape in slide.shapes:
        left, top, right, bottom = _box(shape)
        if left < -_TOLERANCE_IN or top < -_TOLERANCE_IN:
            yield Finding(
                "error",
                "bounds",
                index,
                f"{shape.name} starts off-canvas at ({left:.2f}, {top:.2f})",
            )
        if right > width + _TOLERANCE_IN or bottom > height + _TOLERANCE_IN:
            yield Finding(
                "error",
                "bounds",
                index,
                f"{shape.name} extends to ({right:.2f}, {bottom:.2f}), past {width} x {height}",
            )


def _check_grid(index: int, slide, theme: Theme) -> Iterator[Finding]:
    margin = theme.grid.margin_left_in
    right_edge = theme.canvas.width_in - theme.grid.margin_right_in
    for shape in _generated(slide.shapes):
        left, _, right, _ = _box(shape)
        if margin - _GRID_DRIFT_IN < left < margin - _TOLERANCE_IN:
            yield Finding(
                "error",
                "grid",
                index,
                f"{shape.name} left edge {left:.3f} is inside the margin {margin}",
            )
        if right > right_edge + _TOLERANCE_IN:
            yield Finding(
                "warn",
                "grid",
                index,
                f"{shape.name} right edge {right:.3f} passes {right_edge:.3f}",
            )


def _check_editable(index: int, slide) -> Iterator[Finding]:
    for shape in slide.shapes:
        if shape.shape_type in _RASTER_TYPES:
            yield Finding(
                "error",
                "editable",
                index,
                f"{shape.name} is a picture — diagrams and charts must stay native shapes",
            )


def _check_notes(index: int, slide, theme: Theme) -> Iterator[Finding]:
    text = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
    if len(text) < theme.budget.min_notes_chars:
        yield Finding(
            "error",
            "notes",
            index,
            f"speaker notes are {len(text)} chars, minimum is {theme.budget.min_notes_chars}",
        )
    elif len(text) > theme.budget.max_notes_chars:
        yield Finding(
            "warn",
            "notes",
            index,
            f"speaker notes are {len(text)} chars, over the {theme.budget.max_notes_chars} budget",
        )


def _check_slide_number(index: int, slide) -> Iterator[Finding]:
    if 'type="slidenum"' not in slide._element.xml:
        yield Finding(
            "error",
            "slidenum",
            index,
            "no live slidenum field (masters are not cloned by python-pptx)",
        )


def _check_title(index: int, slide, theme: Theme) -> Iterator[Finding]:
    titles = [s for s in slide.shapes if shape_role(s.name) == "title"]
    for shape in titles:
        text = _all_text(shape).strip()
        if not text:
            continue
        size = theme.typography.title_size(text)
        lines = measure.line_count(
            text,
            width_in=_inches(shape.width),
            size_pt=size,
            typography=theme.typography,
            bold=True,
        )
        if lines > 1:
            yield Finding(
                "error",
                "title",
                index,
                f"title wraps to {lines} lines at {size}pt: {text[:60]!r}",
            )


def _check_text_fit(index: int, slide, theme: Theme) -> Iterator[Finding]:
    for shape in _generated(slide.shapes):
        if not shape.has_text_frame or shape.height is None:
            continue
        frame = shape.text_frame
        if not frame.text.strip():
            continue
        inset = _inches(frame.margin_top) + _inches(frame.margin_bottom)
        usable_w = _inches(shape.width) - _inches(frame.margin_left) - _inches(frame.margin_right)
        total = 0.0
        for paragraph in frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs)
            if not text:
                continue
            size = next(
                (run.font.size.pt for run in paragraph.runs if run.font.size is not None),
                theme.typography.step("body").size_pt,
            )
            bold = any(run.font.bold for run in paragraph.runs)
            lines = measure.line_count(
                text,
                width_in=max(usable_w, 0.2),
                size_pt=size,
                typography=theme.typography,
                bold=bool(bold),
            )
            spacing = paragraph.line_spacing if isinstance(paragraph.line_spacing, float) else 1.1
            total += measure.text_height_in(lines, size_pt=size, line=spacing)
            after = paragraph.space_after
            total += _inches(after) if after is not None else 0.0
        if total > _inches(shape.height) - inset + _TOLERANCE_IN:
            yield Finding(
                "error",
                "overflow",
                index,
                f"{shape.name} needs {total:.2f}in of text in {_inches(shape.height):.2f}in",
            )


def _check_overlap(index: int, slide) -> Iterator[Finding]:
    shapes = _generated(slide.shapes)
    for i, first in enumerate(shapes):
        if shape_role(first.name) in _OVERLAP_OK:
            continue
        for second in shapes[i + 1 :]:
            if shape_role(second.name) in _OVERLAP_OK:
                continue
            a, b = _box(first), _box(second)
            dx = min(a[2], b[2]) - max(a[0], b[0])
            dy = min(a[3], b[3]) - max(a[1], b[1])
            if dx > 0 and dy > 0 and dx * dy > _MIN_OVERLAP_IN2:
                yield Finding(
                    "error",
                    "collision",
                    index,
                    f"{first.name} and {second.name} overlap by {dx * dy:.2f} sq in",
                )


def _check_budget(index: int, slide, theme: Theme) -> Iterator[Finding]:
    for shape in _generated(slide.shapes):
        if shape_role(shape.name) != "bullets" or not shape.has_text_frame:
            continue
        paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
        if len(paragraphs) > theme.budget.max_bullets:
            yield Finding(
                "error",
                "budget",
                index,
                f"{len(paragraphs)} bullets, budget is {theme.budget.max_bullets}",
            )
        for paragraph in paragraphs:
            words = len(paragraph.text.split())
            if words > theme.budget.max_words_per_bullet:
                yield Finding(
                    "error",
                    "budget",
                    index,
                    f"bullet has {words} words, budget is {theme.budget.max_words_per_bullet}: "
                    f"{paragraph.text[:50]!r}",
                )


def _check_font_colors(index: int, slide, theme: Theme) -> Iterator[Finding]:
    allowed = {
        theme.color("body_text"),
        theme.color("secondary_text"),
        theme.color("on_primary"),
        theme.color("page"),
    }
    for shape in _generated(slide.shapes):
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                color = run.font.color
                if color is None or color.type is None:
                    continue
                rgb = getattr(color, "rgb", None)
                if rgb is not None and str(rgb).upper() not in allowed:
                    yield Finding(
                        "warn",
                        "type-colour",
                        index,
                        f"{shape.name} uses font colour #{rgb} — code meaning with fills, not type",
                    )


def _check_takeaway(index: int, slide, theme: Theme) -> Iterator[Finding]:
    roles = {shape_role(s.name) for s in _generated(slide.shapes)}
    titles = [s for s in slide.shapes if shape_role(s.name) == "title"]
    archetype = titles[0].name.split(":")[-1] if titles and titles[0].name.count(":") > 1 else ""
    exempt = {"title", "section", "statement", "agenda", ""}
    if archetype not in exempt and "takeaway" not in roles and theme.chrome.get("takeaway", True):
        yield Finding("warn", "takeaway", index, "no single-takeaway line on this content slide")


def _check_unknowns(index: int, slide) -> Iterator[Finding]:
    text = " ".join(_all_text(s) for s in slide.shapes)
    if not any(token in text.upper() for token in _UNKNOWN_TOKENS):
        return
    upper = text.upper()
    if "OWNER:" not in upper or "DELIVERABLE:" not in upper:
        yield Finding(
            "error",
            "unknowns",
            index,
            "an unknown is marked but no 'Owner:' / 'Deliverable:' pairing appears on the slide",
        )


def _check_smartart(index: int, slide) -> Iterator[Finding]:
    xml = slide._element.xml
    if "/diagram/2006/" in xml or "dgm:relIds" in xml:
        yield Finding("info", "smartart", index, "contains a real SmartArt (dgm) graphic frame")


def _check_balance(index: int, slide, theme: Theme) -> Iterator[Finding]:
    """Catch a content block hanging off the body top with dead space beneath it.

    Slides fail this long before they fail bounds or overlap: everything is legal, the
    grid is honoured, and the slide still reads as unfinished because the whole slack
    landed in one strip above the takeaway.
    """
    body = [
        _box(shape)
        for shape in _generated(slide.shapes)
        if shape_role(shape.name) not in _OUTSIDE_BODY
    ]
    if not body:
        return
    top = min(box[1] for box in body)
    bottom = max(box[3] for box in body)
    above = top - theme.grid.body_top_in
    below = theme.grid.body_bottom_in - bottom
    if below - above > _BALANCE_SLACK_IN:
        yield Finding(
            "warn",
            "balance",
            index,
            f"content leaves {below:.2f} in empty below it against {above:.2f} in above; "
            "place the block with _place() so the slack is shared",
        )


def verify(
    deck_path: Path,
    theme: Theme,
    *,
    do_render: bool = True,
    tier: str = "auto",
    image_dir: Path | None = None,
) -> VerifyReport:
    """Check a built deck and, when possible, rasterise it.

    Args:
        deck_path: The ``.pptx`` to verify.
        theme: Theme whose budgets and grid the deck claims to follow.
        do_render: Run the visual tier as well as the structural checks.
        tier: Force a renderer tier, or ``auto``.
        image_dir: Where slide PNGs are written; defaults to ``<deck>_render/``.

    Returns:
        A :class:`VerifyReport`; ``report.ok`` is False when anything blocking was found.
    """
    presentation = Presentation(str(deck_path))
    report = VerifyReport(deck=deck_path, slide_count=len(presentation.slides._sldIdLst))

    if round(Emu(presentation.slide_width).inches, 2) != round(theme.canvas.width_in, 2) or round(
        Emu(presentation.slide_height).inches, 2
    ) != round(theme.canvas.height_in, 2):
        report.findings.append(
            Finding(
                "error",
                "canvas",
                None,
                f"deck is {Emu(presentation.slide_width).inches:.2f} x "
                f"{Emu(presentation.slide_height).inches:.2f} in, theme expects "
                f"{theme.canvas.width_in} x {theme.canvas.height_in}",
            )
        )

    for number, slide in enumerate(presentation.slides, start=1):
        for check in (
            _check_bounds(number, slide, theme),
            _check_grid(number, slide, theme),
            _check_editable(number, slide),
            _check_notes(number, slide, theme),
            _check_slide_number(number, slide),
            _check_title(number, slide, theme),
            _check_text_fit(number, slide, theme),
            _check_overlap(number, slide),
            _check_budget(number, slide, theme),
            _check_font_colors(number, slide, theme),
            _check_takeaway(number, slide, theme),
            _check_unknowns(number, slide),
            _check_smartart(number, slide),
            _check_balance(number, slide, theme),
        ):
            report.findings.extend(check)

    if do_render:
        target = image_dir or deck_path.with_name(f"{deck_path.stem}_render")
        report.render_result = render(deck_path, target, tier=tier)
        if not report.render_result.ok:
            report.findings.append(
                Finding(
                    "warn",
                    "render",
                    None,
                    f"visual pass skipped ({report.render_result.message}) — "
                    "structural checks only",
                )
            )
    return report
