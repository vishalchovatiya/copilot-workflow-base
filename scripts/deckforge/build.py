"""Render a validated deck spec into a ``.pptx``.

The build owns the parts of a slide that never vary: which layout to use, the title
band, the brand accent bar, the section kicker, the live slide-number field, the
takeaway strip, the unknowns strip and the speaker notes. Everything inside the body
rectangle is delegated to an archetype.

The narrative skeleton is enforced here too — title, agenda, then one divider per
section — because "each slide answers the question the previous one raised" is a
property of slide *order*, not of any single slide.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

from .archetypes import SlideContext, render_slide
from .artifacts import Library
from .primitives import (
    Box,
    TextBlock,
    add_gradient_bar,
    add_shape,
    add_slide_number,
    add_textbox,
    name_shape,
    set_notes,
    write_blocks,
)
from .spec import DeckSpec, SlideSpec
from .tokens import Theme

__all__ = ["BuildError", "BuildResult", "build_deck", "init_template", "resolve_template"]

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_THEME_SLOT_ROLES = (
    ("dk1", "ink"),
    ("lt1", "surface"),
    ("dk2", "primary"),
    ("lt2", "surface_alt"),
    ("accent1", "primary"),
    ("accent2", "secondary"),
    ("accent3", "accent"),
    ("accent4", "caution"),
    ("accent5", "positive"),
    ("accent6", "alert"),
    ("hlink", "primary"),
    ("folHlink", "muted"),
)
_LAYOUT_FOR = {"title": "title", "section": "section", "statement": "blank"}
_SELF_TITLED = frozenset({"title", "statement"})
_MAX_UNKNOWN_LINES = 3


class BuildError(Exception):
    """Raised when a deck cannot be built from the given spec, theme and template."""


@dataclass
class BuildResult:
    """What a build produced."""

    path: Path
    slide_count: int
    template: Path
    theme: str
    warnings: list[str] = field(default_factory=list)
    manifest: list[dict] = field(default_factory=list)

    def manifest_path(self) -> Path:
        """Where the per-slide build manifest was written."""
        return self.path.with_suffix(".build.json")


def resolve_template(theme: Theme, spec: DeckSpec | None, search: list[Path]) -> Path:
    """Locate the template this deck must be built from.

    Building from the stock Office template is never correct for a real deck: it has no
    branding, no corporate fonts and no logo. So the template is resolved explicitly and
    a missing one is an error with a fix attached, not a silent fallback.
    """
    candidates = [spec.template] if spec and spec.template else []
    if theme.template:
        candidates.append(theme.template)
    for candidate in candidates:
        direct = Path(candidate)
        if direct.is_file():
            return direct
        for folder in search:
            found = folder / candidate
            if found.is_file():
                return found
    wanted = candidates[0] if candidates else "<none configured>"
    raise BuildError(
        f"template {wanted!r} not found in {', '.join(str(p) for p in search)}. "
        "Harvest one from a branded deck with `deckforge harvest <deck.pptx>`, or "
        "generate the neutral base with `deckforge init-template`."
    )


def init_template(theme: Theme, destination: Path) -> Path:
    """Manufacture the neutral base template from the theme tokens.

    This is the only place in the package that opens the stock Office template, and it
    does so to *build* a template, never to build a deck: the slide size, colour scheme
    and font scheme are rewritten from tokens so the result is a real, self-consistent
    base package.
    """
    presentation = Presentation()
    presentation.slide_width = Inches(theme.canvas.width_in)
    presentation.slide_height = Inches(theme.canvas.height_in)
    part = presentation.slide_masters[0].part.part_related_by(RT.THEME)
    tree = etree.fromstring(part.blob)
    scheme = tree.find(f".//{_A}clrScheme")
    if scheme is None:
        raise BuildError("stock template has no colour scheme to rewrite")
    for slot, role in _THEME_SLOT_ROLES:
        node = scheme.find(f"{_A}{slot}")
        if node is None:
            continue
        for child in list(node):
            node.remove(child)
        colour = etree.SubElement(node, f"{_A}srgbClr")
        colour.set("val", theme.color(role))
    for tag, key in ((f"{_A}majorFont", "title"), (f"{_A}minorFont", "body")):
        latin = tree.find(f".//{tag}/{_A}latin")
        if latin is not None:
            latin.set("typeface", theme.typography.fonts[key])
    part._blob = etree.tostring(tree, xml_declaration=True, encoding="UTF-8")
    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(destination))
    return destination


def _strip_placeholders(slide, keep_title: bool) -> None:
    """Drop unused inherited placeholders so no 'Click to add' prompt survives."""
    for shape in list(slide.placeholders):
        if keep_title and shape == slide.shapes.title:
            continue
        shape._element.getparent().remove(shape._element)


def _add_title(slide, theme: Theme, spec: SlideSpec) -> None:
    grid = theme.grid
    box = Box(grid.margin_left_in, grid.title_top_in, theme.title_width, grid.title_height_in)
    size = theme.typography.title_size(spec.title)
    block = TextBlock(text=spec.title, step="title")
    placeholder = slide.shapes.title
    if placeholder is not None:
        placeholder.name = name_shape("title", spec.archetype)
        placeholder.left, placeholder.top = Inches(box.left), Inches(box.top)
        placeholder.width, placeholder.height = Inches(box.width), Inches(box.height)
        placeholder.text_frame.word_wrap = True
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(placeholder.text_frame, side, Inches(0))
        write_blocks(placeholder, theme, [block])
        shape = placeholder
    else:
        shape = add_textbox(
            slide, box, theme, [block], role="title", detail=spec.archetype, anchor="bottom"
        )
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)


def _add_chrome(slide, theme: Theme, spec: SlideSpec) -> None:
    grid = theme.grid
    if theme.chrome.get("accent_bar", True):
        add_gradient_bar(
            slide,
            Box(
                grid.margin_left_in, grid.accent_top_in, grid.accent_width_in, grid.accent_height_in
            ),
            theme,
            "primary",
            "secondary",
        )
    if theme.chrome.get("kicker", True) and spec.section and spec.archetype != "section":
        left = grid.margin_left_in + grid.accent_width_in + 0.3
        add_textbox(
            slide,
            # The kicker honours the same right-hand guard as the title, so it cannot
            # land under a master logo.
            Box(
                left,
                grid.accent_top_in - 0.10,
                theme.title_width - left + grid.margin_left_in,
                0.26,
            ),
            theme,
            [TextBlock(text=spec.section, step="eyebrow", color="secondary_text", align="right")],
            role="kicker",
            anchor="middle",
        )


def _add_takeaway(slide, theme: Theme, spec: SlideSpec, role: str, bottom: float) -> float:
    """Draw the slide's single claim; return the height it consumed."""
    if not spec.takeaway or not theme.chrome.get("takeaway", True):
        return 0.0
    grid, spacing = theme.grid, theme.spacing
    height = spacing.callout_height_in
    box = Box(grid.margin_left_in, bottom - height, theme.content_width, height)
    add_shape(
        slide,
        MSO_SHAPE.RECTANGLE,
        box,
        role="takeaway",
        fill=theme.tint(role, "wash"),
        line=None,
    )
    add_shape(
        slide,
        MSO_SHAPE.RECTANGLE,
        Box(box.left, box.top, spacing.accent_edge_in, height),
        role="band",
        detail="takeaway",
        fill=theme.color(role),
    )
    add_textbox(
        slide,
        Box(box.left + spacing.card_pad_in, box.top, box.width - spacing.card_pad_in * 2, height),
        theme,
        [TextBlock(text=spec.takeaway, step="small", bold=True)],
        role="label",
        detail="takeaway",
        anchor="middle",
    )
    return height + theme.spacing.row_gap_in


def _add_unknowns(slide, theme: Theme, spec: SlideSpec, bottom: float) -> float:
    """Render unknowns as explicit placeholders with an owner and a deliverable."""
    if not spec.unknowns:
        return 0.0
    shown = spec.unknowns[:_MAX_UNKNOWN_LINES]
    line_height = 0.24
    height = line_height * len(shown)
    box = Box(theme.grid.margin_left_in, bottom - height, theme.content_width, height)
    add_textbox(
        slide,
        box,
        theme,
        [TextBlock(text=unknown.label(), step="tiny", color="secondary_text") for unknown in shown],
        role="unknowns",
    )
    return height + theme.spacing.row_gap_in


def _notes_text(spec: SlideSpec) -> str:
    extra = "".join(f"\nUnknown — {unknown.label()}" for unknown in spec.unknowns)
    return f"{spec.notes.strip()}{extra}"


def _skeleton(deck: DeckSpec) -> list[SlideSpec]:
    """Expand the authored spec into the full slide order, dividers and all."""
    slides: list[SlideSpec] = []
    authored = list(deck.slides())
    has_title = bool(authored) and authored[0].archetype == "title"
    if not has_title:
        slides.append(
            SlideSpec(
                archetype="title",
                title=deck.title,
                notes=(
                    f"Opening slide for {deck.title}. State the promise of the deck in one "
                    "sentence, then move straight to the agenda: the audience should know "
                    "what decision they are being asked for before any mechanism appears."
                ),
                content={"subtitle": deck.subtitle},
            )
        )
    if deck.agenda:
        slides.append(
            SlideSpec(
                archetype="agenda",
                title="Agenda",
                notes=(
                    "Walk the agenda once and name the question each section answers. "
                    "This is the contract for the rest of the deck: every later divider "
                    "shows progress against exactly this list, so the audience always "
                    "knows which question is currently open."
                ),
                content={},
            )
        )
    for index, section in enumerate(deck.sections, start=1):
        if len(deck.sections) > 1:
            slides.append(
                SlideSpec(
                    archetype="section",
                    title=section.name,
                    notes=(
                        f"Divider for {section.name}. "
                        + (
                            f"The question this section answers: {section.question} "
                            if section.question
                            else ""
                        )
                        + "Pause here, restate the previous section's answer in one line, "
                        "then open this question so the deck reads as one continuous "
                        "argument rather than a pile of slides."
                    ),
                    content={"question": section.question},
                    section=section.name,
                    section_index=index,
                )
            )
        slides.extend(section.slides)
    return slides


def build_deck(
    spec: DeckSpec,
    theme: Theme,
    output: Path,
    *,
    template: Path,
    library: Library | None = None,
) -> BuildResult:
    """Render a spec into a fully-editable ``.pptx``.

    Args:
        spec: Validated deck spec.
        theme: Active theme; supplies every colour, coordinate and type size.
        output: Destination ``.pptx``.
        template: Base package to build on — harvested or generated, never implicit.
        library: Artifact library, required only if a slide uses the artifact archetype.

    Returns:
        A :class:`BuildResult` including the per-slide manifest written next to the deck.
    """
    presentation = Presentation(str(template))
    presentation.slide_width = Inches(theme.canvas.width_in)
    presentation.slide_height = Inches(theme.canvas.height_in)
    slide_ids = presentation.slides._sldIdLst
    for slide_id in list(slide_ids):
        presentation.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)

    result = BuildResult(
        path=output,
        slide_count=0,
        template=template,
        theme=theme.name,
        warnings=list(spec.warnings),
    )
    layouts = presentation.slide_layouts
    for number, slide_spec in enumerate(_skeleton(spec), start=1):
        layout_role = _LAYOUT_FOR.get(slide_spec.archetype, "content")
        index = min(theme.layout_index(layout_role), len(layouts) - 1)
        slide = presentation.slides.add_slide(layouts[index])
        self_titled = slide_spec.archetype in _SELF_TITLED
        _strip_placeholders(slide, keep_title=not self_titled)
        role = spec.entities.get(slide_spec.entity, "primary")

        if not self_titled:
            _add_title(slide, theme, slide_spec)
            _add_chrome(slide, theme, slide_spec)
        bottom = theme.grid.body_bottom_in
        bottom -= _add_takeaway(slide, theme, slide_spec, role, bottom)
        bottom -= _add_unknowns(slide, theme, slide_spec, bottom)
        body = Box(
            theme.grid.margin_left_in,
            theme.grid.body_top_in,
            theme.content_width,
            bottom - theme.grid.body_top_in,
        )
        render_slide(
            SlideContext(
                slide=slide, theme=theme, deck=spec, spec=slide_spec, body=body, library=library
            )
        )
        if theme.chrome.get("slide_number", True):
            add_slide_number(
                slide,
                theme,
                color="on_primary" if slide_spec.archetype == "statement" else "secondary_text",
            )
        set_notes(slide, _notes_text(slide_spec))
        result.manifest.append(
            {
                "slide": number,
                "archetype": slide_spec.archetype,
                "section": slide_spec.section,
                "title": slide_spec.title,
                "takeaway": slide_spec.takeaway,
                "entity": slide_spec.entity,
                "entity_role": role if slide_spec.entity else "",
                "artifact": slide_spec.content.get("artifact", ""),
                "notes_chars": len(_notes_text(slide_spec)),
                "unknowns": [unknown.label() for unknown in slide_spec.unknowns],
            }
        )
        result.slide_count = number

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    result.manifest_path().write_text(
        json.dumps(
            {
                "deck": output.name,
                "built": datetime.now(timezone.utc).date().isoformat(),
                "theme": theme.name,
                "theme_source": str(theme.source),
                "template": str(template),
                "spec": str(spec.source),
                "artifact_library": str(library.root) if library else "",
                "warnings": result.warnings,
                "slides": result.manifest,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    return result
