"""Slide archetypes: one function per visual pattern, all of them grid-driven.

An archetype receives a body rectangle and the slide's content, and draws native
shapes into it. It never chooses a colour or a coordinate of its own — both come from
the theme — which is what lets the same spec render on-brand against a second theme.

Diagrams here are rebuilt from autoshapes rather than emitted as SmartArt. python-pptx
cannot create a ``dgm`` graphic frame, and a rebuilt diagram is fully recolourable from
tokens, so this is the honest default. Genuine SmartArt reaches a deck only through the
``artifact`` archetype, which clones a harvested frame together with its diagram parts.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass

from pptx.enum.shapes import MSO_SHAPE

from . import measure
from .artifacts import Library
from .primitives import (
    Box,
    TextBlock,
    add_bullets,
    add_connector,
    add_shape,
    add_table,
    add_textbox,
    set_cell_borders,
    style_cell,
)
from .spec import DeckSpec, SlideSpec
from .tokens import Theme

__all__ = ["ARCHETYPE_RENDERERS", "SlideContext", "render_slide"]

_TICK = "\u2713"
_CROSS = "\u2717"


@dataclass
class SlideContext:
    """Everything an archetype needs, and nothing it should not have."""

    slide: object
    theme: Theme
    deck: DeckSpec
    spec: SlideSpec
    body: Box
    library: Library | None = None

    @property
    def role(self) -> str:
        """Semantic colour role for this slide's entity, or the deck default."""
        return self.deck.entities.get(self.spec.entity, "primary")

    def fill(self, level: str = "panel") -> str:
        """Tinted fill for this slide's entity colour."""
        return self.theme.tint(self.role, level)

    def edge(self) -> str:
        """Full-strength entity colour, for bars, borders and connectors."""
        return self.theme.color(self.role)


def _block_height(theme: Theme, text: str, width: float, step_name: str) -> float:
    """Height one paragraph occupies at a given width, including its trailing space."""
    if not text:
        return 0.0
    step = theme.typography.step(step_name)
    lines = measure.line_count(
        text,
        width_in=max(width, 0.2),
        size_pt=step.size_pt,
        typography=theme.typography,
        bold=step.bold,
        caps=step.caps,
    )
    return (
        measure.text_height_in(lines, size_pt=step.size_pt, line=step.line)
        + step.space_after_pt / 72.0
    )


def _place(theme: Theme, body: Box, height: float) -> Box:
    """Sit a block in the body band, splitting the leftover height above and below it.

    Hanging every block off the body top dumps the whole slack into one dead strip
    above the callout, which is the single most common way a generated slide looks
    unfinished. ``layout.block_bias`` decides how much of it goes above.
    """
    used = min(height, body.height)
    top = body.top + max(0.0, (body.height - used) * theme.layout.block_bias)
    return Box(body.left, top, body.width, used)


def _fill_floor(ctx: SlideContext, rows: int = 1) -> float:
    """Least height one box in a row should take, as a share of the body band.

    Sizing a box to its own copy alone makes a two-line card render as a thin strip
    stranded in white space; the floor keeps the block in proportion to the slide.
    """
    gap = ctx.theme.spacing.card_gap_in
    span = ctx.body.height * ctx.theme.layout.block_fill - gap * (rows - 1)
    return max(0.0, span) / max(rows, 1)


def _centre_stack(inner: Box, content_height: float) -> Box:
    """Centre a content stack inside a box grown past it, so growth reads as padding."""
    slack = max(0.0, inner.height - content_height)
    return Box(inner.left, inner.top + slack / 2, inner.width, inner.height - slack / 2)


def _card_height(
    ctx: SlideContext,
    width: float,
    label: str,
    text: str,
    *,
    label_height: float = 0.0,
    floor: float = 0.0,
) -> float:
    """How tall a card has to be for its own content, within the cap."""
    pad = ctx.theme.spacing.card_pad_in
    inner = width - 2 * pad
    layout = ctx.theme.layout
    needed = (
        2 * pad
        + (label_height or _block_height(ctx.theme, label, inner, "small"))
        + _block_height(ctx.theme, text, inner, "micro")
    )
    height = max(layout.card_min_height_in, min(needed, layout.card_max_height_in))
    return max(height, min(floor, layout.card_max_height_in))


@dataclass(frozen=True)
class CardRow:
    """Measurements one row of cards shares so the row reads as one band."""

    label_height: float
    height: float
    content_height: float


def _card_row(
    ctx: SlideContext, width: float, entries: list[tuple[str, str]], *, rows: int = 1
) -> CardRow:
    """Measure one row of cards once, so every card in it lines up.

    Label height, card height and content height are all shared: a label wrapping to
    two lines in one card must not push that card's body text out of line with its
    neighbours, and neither must a longer sentence in the card next to it.
    """
    pad = ctx.theme.spacing.card_pad_in
    inner = width - 2 * pad
    floor = _fill_floor(ctx, rows)
    label_height = max(_block_height(ctx.theme, label, inner, "small") for label, _ in entries)
    height = max(
        _card_height(ctx, width, label, text, label_height=label_height, floor=floor)
        for label, text in entries
    )
    content = label_height + max(
        _block_height(ctx.theme, text, inner, "micro") for _, text in entries
    )
    return CardRow(label_height, height, content)


def _card(
    ctx: SlideContext,
    box: Box,
    label: str,
    text: str,
    *,
    index: int,
    label_height: float = 0.0,
    content_height: float = 0.0,
) -> None:
    """One tinted card with a full-strength top edge and two type sizes."""
    theme = ctx.theme
    add_shape(
        ctx.slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        box,
        role="card",
        detail=str(index),
        fill=ctx.fill(),
        line=theme.color("hairline"),
        adjustment=theme.layout.card_round_adjust,
    )
    add_shape(
        ctx.slide,
        MSO_SHAPE.RECTANGLE,
        Box(box.left, box.top, box.width, theme.spacing.accent_edge_in),
        role="band",
        detail=str(index),
        fill=ctx.edge(),
    )
    pad = theme.spacing.card_pad_in
    inner = Box(box.left + pad, box.top + pad, box.width - 2 * pad, box.height - 2 * pad)
    head = label_height or _block_height(theme, label, inner.width, "small")
    body_height = _block_height(theme, text, inner.width, "micro")
    inner = _centre_stack(inner, content_height or head + body_height)
    add_textbox(
        ctx.slide,
        inner.top_slice(head),
        theme,
        [TextBlock(text=label, step="small", bold=True)],
        role="label",
        detail=str(index),
    )
    if text:
        add_textbox(
            ctx.slide,
            Box(inner.left, inner.top + head, inner.width, max(inner.height - head, 0.2)),
            theme,
            [TextBlock(text=text, step="micro", color="secondary_text")],
            role="label",
            detail=f"body{index}",
        )


def _panel_content(ctx: SlideContext, width: float, heading: str, items: list[str]) -> float:
    """Height of a panel's heading-plus-bullets stack, excluding its padding."""
    theme = ctx.theme
    pad = theme.spacing.card_pad_in
    inner = width - 2 * pad - theme.spacing.bullet_indent_in
    return (
        _block_height(theme, heading, width - 2 * pad, "lead")
        + theme.layout.head_gap_in
        + sum(_block_height(theme, item, inner, "micro") for item in items)
    )


def _panel_height(
    ctx: SlideContext, width: float, heading: str, items: list[str], *, floor: float = 0.0
) -> float:
    """How tall a panel has to be for its heading and bullets."""
    theme = ctx.theme
    needed = theme.spacing.card_pad_in * 2.4 + _panel_content(ctx, width, heading, items)
    height = max(theme.layout.card_min_height_in, min(needed, theme.grid.body_height))
    return max(height, min(floor or _fill_floor(ctx), theme.grid.body_height))


def _panel(
    ctx: SlideContext,
    box: Box,
    heading: str,
    items: list[str],
    role: str,
    index: int,
    *,
    bullet: str = "\u2013",
    content_height: float = 0.0,
) -> None:
    theme = ctx.theme
    add_shape(
        ctx.slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        box,
        role="card",
        detail=f"panel{index}",
        fill=theme.tint(role, "panel"),
        line=theme.color("hairline"),
        adjustment=theme.layout.card_round_adjust,
    )
    add_shape(
        ctx.slide,
        MSO_SHAPE.RECTANGLE,
        Box(box.left, box.top, box.width, theme.spacing.accent_edge_in),
        role="band",
        detail=f"panel{index}",
        fill=theme.color(role),
    )
    pad = theme.spacing.card_pad_in
    inner = Box(box.left + pad, box.top + pad * 1.4, box.width - 2 * pad, box.height - pad * 2.4)
    head_height = _block_height(theme, heading, inner.width, "lead")
    items_height = sum(
        _block_height(theme, item, inner.width - theme.spacing.bullet_indent_in, "micro")
        for item in items
    )
    inner = _centre_stack(
        inner, content_height or head_height + theme.layout.head_gap_in + items_height
    )
    add_textbox(
        ctx.slide,
        inner.top_slice(head_height),
        theme,
        [TextBlock(text=heading, step="lead", bold=True)],
        role="label",
        detail=f"panelhead{index}",
    )
    add_bullets(
        ctx.slide,
        Box(
            inner.left,
            inner.top + head_height + theme.layout.head_gap_in,
            inner.width,
            inner.height - head_height - theme.layout.head_gap_in,
        ),
        theme,
        items,
        step="micro",
        bullet=bullet,
        bullet_color=role,
        role="label",
    )


def archetype_title(ctx: SlideContext) -> None:
    """Slide 1: the promise of the deck, nothing else."""
    theme, deck = ctx.theme, ctx.deck
    grid, layout = theme.grid, theme.layout
    band = Box(
        grid.margin_left_in,
        layout.hero_top_in,
        theme.content_width * layout.hero_width_ratio,
        layout.hero_height_in,
    )
    add_textbox(
        ctx.slide,
        band,
        theme,
        [TextBlock(text=deck.title, step="hero")],
        role="title",
        detail="title",
        anchor="bottom",
    )
    add_shape(
        ctx.slide,
        MSO_SHAPE.RECTANGLE,
        Box(
            grid.margin_left_in,
            band.bottom + layout.hero_rule_gap_in,
            layout.hero_rule_width_in,
            grid.accent_height_in,
        ),
        role="accent",
        fill=theme.color("primary"),
    )
    subtitle = ctx.spec.content.get("subtitle", deck.subtitle)
    footer = " \u00b7 ".join(
        part
        for part in (
            ctx.spec.content.get("presenter", deck.presenter),
            ctx.spec.content.get("date", deck.date),
        )
        if part
    )
    blocks = (
        [TextBlock(text=subtitle, step="lead", color="secondary_text", space_after_pt=10)]
        if subtitle
        else []
    )
    if footer:
        blocks.append(TextBlock(text=footer, step="eyebrow", color="secondary_text"))
    if blocks:
        add_textbox(
            ctx.slide,
            Box(
                grid.margin_left_in,
                band.bottom + layout.hero_sub_gap_in,
                theme.content_width * layout.hero_width_ratio,
                layout.hero_height_in,
            ),
            theme,
            blocks,
            role="label",
            detail="subtitle",
        )


def archetype_agenda(ctx: SlideContext) -> None:
    """The explicit contract: every section, in order, with the question it answers."""
    theme = ctx.theme
    entries = ctx.spec.content.get("items") or [
        {"label": section.name, "text": section.question} for section in ctx.deck.sections
    ]
    columns = min(len(entries), 4)
    row_count = -(-len(entries) // columns)
    gap = theme.spacing.card_gap_in
    width = (ctx.body.width - gap * (columns - 1)) / columns
    labelled = [
        (f"{number:02d}   {entry.get('label', '')}", entry.get("text", ""))
        for number, entry in enumerate(entries, start=1)
    ]
    rows = [labelled[row * columns : (row + 1) * columns] for row in range(row_count)]
    measured = [_card_row(ctx, width, chunk, rows=row_count) for chunk in rows]
    block = _place(theme, ctx.body, sum(row.height for row in measured) + gap * (row_count - 1))
    top = block.top
    for row_index, (chunk, row) in enumerate(zip(rows, measured)):
        for offset, (label, text) in enumerate(chunk):
            _card(
                ctx,
                Box(ctx.body.left + offset * (width + gap), top, width, row.height),
                label,
                text,
                index=row_index * columns + offset + 1,
                label_height=row.label_height,
                content_height=row.content_height,
            )
        top += row.height + gap


def archetype_section(ctx: SlideContext) -> None:
    """A divider that also shows where we are in the agenda.

    The section name is already in the title band, so the body carries only the numeral,
    the question this section answers and the progress strip.
    """
    theme = ctx.theme
    layout = theme.layout
    index = ctx.spec.section_index
    strip_height = layout.progress_height_in + layout.progress_label_gap_in
    band = ctx.body.shrink_bottom(strip_height + theme.spacing.panel_gap_in)
    row = _place(theme, band, layout.numeral_height_in)
    numeral = Box(row.left, row.top, layout.numeral_width_in, row.height)
    add_textbox(
        ctx.slide,
        numeral,
        theme,
        [TextBlock(text=f"{index:02d}", step="hero", color="secondary_text")],
        role="label",
        detail="numeral",
        anchor="middle",
    )
    question = ctx.spec.content.get("question", "")
    if question:
        add_textbox(
            ctx.slide,
            Box(
                numeral.right + layout.numeral_gap_in,
                numeral.top,
                row.width - numeral.width - layout.numeral_gap_in,
                numeral.height,
            ),
            theme,
            [TextBlock(text=question, step="lead", color="secondary_text")],
            role="label",
            detail="question",
            anchor="middle",
        )

    progress = Box(ctx.body.left, ctx.body.bottom - strip_height, ctx.body.width, strip_height)
    lanes = progress.columns(max(len(ctx.deck.sections), 1), theme.spacing.card_gap_in)
    for position, (lane, section) in enumerate(zip(lanes, ctx.deck.sections), start=1):
        done = position <= index
        add_shape(
            ctx.slide,
            MSO_SHAPE.RECTANGLE,
            Box(lane.left, lane.top, lane.width, theme.spacing.accent_edge_in),
            role="band",
            detail=f"progress{position}",
            fill=theme.color("primary") if done else theme.color("hairline"),
        )
        add_textbox(
            ctx.slide,
            Box(
                lane.left,
                lane.top + layout.progress_label_gap_in,
                lane.width,
                layout.progress_height_in,
            ),
            theme,
            [
                TextBlock(
                    text=section.name,
                    step="tiny",
                    color="body_text" if position == index else "secondary_text",
                )
            ],
            role="label",
            detail=f"progresslabel{position}",
        )


def archetype_bullets(ctx: SlideContext) -> None:
    """The default content slide: one lead line, then short bullets."""
    theme = ctx.theme
    lead = ctx.spec.content.get("lead", "")
    items = list(ctx.spec.content["bullets"])
    gap = theme.layout.lead_gap_in
    lead_height = (_block_height(theme, lead, ctx.body.width, "lead") + gap) if lead else 0.0
    bullet_width = ctx.body.width - theme.spacing.bullet_indent_in
    body_height = sum(_block_height(theme, item, bullet_width, "body") for item in items)
    block = _place(theme, ctx.body, lead_height + body_height)
    if lead:
        add_textbox(
            ctx.slide,
            block.top_slice(lead_height - gap),
            theme,
            [TextBlock(text=lead, step="lead", color="secondary_text")],
            role="label",
            detail="lead",
        )
    add_bullets(
        ctx.slide,
        Box(
            block.left,
            block.top + lead_height,
            block.width,
            max(block.height - lead_height, 0.2),
        ),
        theme,
        items,
        step="body",
        bullet_color=ctx.role,
        anchor="top",
    )


def archetype_cards(ctx: SlideContext) -> None:
    """Two to five parallel ideas, one card each."""
    cards = ctx.spec.content["cards"]
    columns = ctx.body.columns(len(cards), ctx.theme.spacing.card_gap_in)
    entries = [(card.get("label", ""), card.get("text", "")) for card in cards]
    row = _card_row(ctx, columns[0].width, entries)
    block = _place(ctx.theme, ctx.body, row.height)
    for index, ((label, text), target) in enumerate(zip(entries, columns), start=1):
        _card(
            ctx,
            Box(target.left, block.top, target.width, block.height),
            label,
            text,
            index=index,
            label_height=row.label_height,
            content_height=row.content_height,
        )


def archetype_process(ctx: SlideContext) -> None:
    """A left-to-right flow as a chevron chain, rebuilt from autoshapes."""
    theme = ctx.theme
    layout = theme.layout
    steps = ctx.spec.content["steps"]
    bullets = list(ctx.spec.content.get("bullets") or [])
    captions = [step.get("text", "") for step in steps]
    caption_height = layout.caption_gap_in + layout.caption_height_in if any(captions) else 0.0
    bullet_width = ctx.body.width - theme.spacing.bullet_indent_in
    bullet_height = sum(_block_height(theme, item, bullet_width, "small") for item in bullets)
    if bullets:
        bullet_height += theme.spacing.panel_gap_in
    block = _place(theme, ctx.body, layout.chevron_height_in + caption_height + bullet_height)
    chain = Box(block.left, block.top, block.width, layout.chevron_height_in)
    overlap = theme.spacing.chevron_overlap_in
    width = (chain.width + overlap * (len(steps) - 1)) / len(steps)
    inset = theme.spacing.chevron_text_inset_in
    label_pad = theme.spacing.row_gap_in
    for index, step in enumerate(steps):
        left = chain.left + index * (width - overlap)
        shade = "mid" if index % 2 else "chip"
        add_shape(
            ctx.slide,
            MSO_SHAPE.CHEVRON,
            Box(left, chain.top, width, chain.height),
            role="chevron",
            detail=str(index + 1),
            fill=theme.tint(ctx.role, shade),
            line=None,
            adjustment=layout.chevron_adjust,
        )
        add_textbox(
            ctx.slide,
            Box(
                left + inset,
                chain.top + label_pad,
                width - 2 * inset,
                chain.height - 2 * label_pad,
            ),
            theme,
            [TextBlock(text=step.get("label", ""), step="small", bold=True, align="center")],
            role="label",
            detail=f"step{index + 1}",
            anchor="middle",
        )
    if any(captions):
        caption_box = Box(
            chain.left,
            chain.bottom + layout.caption_gap_in,
            chain.width,
            layout.caption_height_in,
        )
        for index, (caption, target) in enumerate(
            zip(captions, caption_box.columns(len(steps), theme.spacing.card_gap_in)), start=1
        ):
            if not caption:
                continue
            add_textbox(
                ctx.slide,
                target,
                theme,
                [TextBlock(text=caption, step="micro", color="secondary_text", align="center")],
                role="label",
                detail=f"caption{index}",
            )
    if bullets:
        top = chain.bottom + caption_height + theme.spacing.panel_gap_in
        add_bullets(
            ctx.slide,
            Box(block.left, top, block.width, max(block.bottom - top, 0.2)),
            theme,
            bullets,
            step="small",
            bullet_color=ctx.role,
            anchor="top",
        )


def archetype_panels(ctx: SlideContext) -> None:
    """A versus B, colour-coded, never more than two."""
    theme = ctx.theme
    panels = ctx.spec.content["panels"]
    roles = [ctx.role, _contrast_role(ctx)]
    columns = ctx.body.columns(2, theme.spacing.panel_gap_in)
    height = max(
        _panel_height(
            ctx, columns[0].width, panel.get("heading", ""), list(panel.get("bullets", []))
        )
        for panel in panels
    )
    content = max(
        _panel_content(
            ctx, columns[0].width, panel.get("heading", ""), list(panel.get("bullets", []))
        )
        for panel in panels
    )
    block = _place(theme, ctx.body, height)
    for index, (panel, target) in enumerate(zip(panels, columns), start=1):
        _panel(
            ctx,
            Box(target.left, block.top, target.width, block.height),
            panel.get("heading", ""),
            list(panel.get("bullets", [])),
            roles[index - 1],
            index,
            content_height=content,
        )


def _contrast_role(ctx: SlideContext) -> str:
    roles = list(ctx.theme.entity_roles) or ["primary"]
    current = ctx.role
    return next((role for role in roles if role != current), current)


def archetype_checklist(ctx: SlideContext) -> None:
    """Do and do-not, with glyphs instead of colour-coded type."""
    theme = ctx.theme
    content = ctx.spec.content
    columns = ctx.body.columns(2, theme.spacing.panel_gap_in)
    pairs = (
        (content["do"], "positive", _TICK, 1),
        (content["dont"], "alert", _CROSS, 2),
    )
    marked = [
        [f"{glyph}  {item}" for item in payload.get("items", [])] for payload, _, glyph, _ in pairs
    ]
    height = max(
        _panel_height(ctx, columns[0].width, payload.get("heading", ""), items)
        for (payload, _, _, _), items in zip(pairs, marked)
    )
    content_height = max(
        _panel_content(ctx, columns[0].width, payload.get("heading", ""), items)
        for (payload, _, _, _), items in zip(pairs, marked)
    )
    block = _place(theme, ctx.body, height)
    for ((payload, role, _, index), items), target in zip(zip(pairs, marked), columns):
        _panel(
            ctx,
            Box(target.left, block.top, target.width, block.height),
            payload.get("heading", ""),
            items,
            role,
            index,
            bullet="",
            content_height=content_height,
        )


def archetype_table(ctx: SlideContext) -> None:
    """A native table: horizontal rules only, first column carrying the entity tint."""
    theme = ctx.theme
    columns = ctx.spec.content["columns"]
    rows = ctx.spec.content["rows"]
    legend = ctx.spec.content.get("legend", "")
    layout = theme.layout
    reserved = (layout.legend_gap_in + layout.legend_height_in) if legend else 0.0
    height = min(ctx.body.height - reserved, layout.table_max_height_in)
    box = Box(ctx.body.left, ctx.body.top, ctx.body.width, height)
    step = "micro" if len(columns) > 5 else "small"
    _, table = add_table(
        ctx.slide, box, theme, rows=len(rows) + 1, cols=len(columns), detail="main"
    )
    rule = theme.color("body_text")
    hairline = theme.color("hairline")
    for index, heading in enumerate(columns):
        cell = table.cell(0, index)
        style_cell(cell, theme, str(heading), step=step, bold=True, fill=theme.color("page"))
        set_cell_borders(cell, top=rule, bottom=rule, width_pt=theme.spacing.hairline_pt)
    for row_index, row in enumerate(rows, start=1):
        last = row_index == len(rows)
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            fill = theme.color("page") if row_index % 2 else theme.tint(ctx.role, "wash")
            if col_index == 0:
                fill = theme.tint(ctx.role, "chip")
            style_cell(
                cell,
                theme,
                str(value),
                step=step,
                fill=fill,
                bold=col_index == 0,
                align="left",
            )
            set_cell_borders(
                cell,
                bottom=rule if last else hairline,
                width_pt=theme.spacing.hairline_pt,
            )
    if legend:
        add_textbox(
            ctx.slide,
            Box(box.left, box.bottom + layout.legend_gap_in, box.width, layout.legend_height_in),
            theme,
            [TextBlock(text=legend, step="tiny", color="secondary_text")],
            role="label",
            detail="legend",
        )


def archetype_artifact(ctx: SlideContext) -> None:
    """Re-place a harvested diagram, table or SmartArt from the artifact library."""
    if ctx.library is None:
        raise ValueError(
            f"slide {ctx.spec.title!r} uses the artifact archetype but no library is "
            "configured; set `deck.artifacts:` in the spec or pass --artifacts"
        )
    theme = ctx.theme
    layout = theme.layout
    content = ctx.spec.content
    caption = content.get("caption", "")
    bullets = content.get("bullets") or []
    artifact = ctx.library.get(content["artifact"])
    # A wide diagram squeezed into half the slide reads as a thumbnail; give it the
    # full content width and let its bullets sit underneath instead.
    stacked = bool(bullets) and artifact.aspect >= layout.artifact_wide_ratio
    box = ctx.body
    reserve = 0.0
    if bullets and not stacked:
        box = Box(
            ctx.body.left,
            ctx.body.top,
            ctx.body.width * layout.artifact_split_ratio,
            ctx.body.height,
        )
        side = Box(
            box.right + theme.spacing.panel_gap_in,
            ctx.body.top,
            ctx.body.width - box.width - theme.spacing.panel_gap_in,
            ctx.body.height,
        )
        add_bullets(
            ctx.slide, side, theme, bullets, step="small", bullet_color=ctx.role, anchor="middle"
        )
    elif stacked:
        reserve = sum(_block_height(theme, item, ctx.body.width, "small") for item in bullets)
        box = box.shrink_bottom(reserve + theme.spacing.panel_gap_in)
    if caption:
        box = box.shrink_bottom(layout.caption_gap_in + layout.legend_height_in)
    placed = ctx.library.place(ctx.slide, artifact, theme, box, text=content.get("text"))
    bottom = placed.bottom
    if caption:
        add_textbox(
            ctx.slide,
            Box(
                box.left,
                placed.bottom + layout.caption_gap_in,
                box.width,
                layout.legend_height_in,
            ),
            theme,
            [TextBlock(text=caption, step="tiny", color="secondary_text")],
            role="label",
            detail="caption",
        )
        bottom += layout.caption_gap_in + layout.legend_height_in
    if stacked:
        add_bullets(
            ctx.slide,
            Box(ctx.body.left, bottom + theme.spacing.panel_gap_in, ctx.body.width, reserve),
            theme,
            bullets,
            step="small",
            bullet_color=ctx.role,
            anchor="top",
        )


def archetype_statement(ctx: SlideContext) -> None:
    """One line, full bleed, white on brand colour. Closing only."""
    theme = ctx.theme
    layout = theme.layout
    canvas = Box(0, 0, theme.canvas.width_in, theme.canvas.height_in)
    add_shape(ctx.slide, MSO_SHAPE.RECTANGLE, canvas, role="backdrop", fill=theme.color("primary"))
    text = ctx.spec.content["statement"]
    add_textbox(
        ctx.slide,
        Box(
            theme.grid.margin_left_in,
            layout.hero_top_in,
            theme.content_width * layout.statement_width_ratio,
            theme.canvas.height_in - 2 * layout.hero_top_in,
        ),
        theme,
        [TextBlock(text=text, step="hero", color="on_primary")],
        role="title",
        detail="statement",
        anchor="middle",
    )


def archetype_summary(ctx: SlideContext) -> None:
    """The closing picture: the whole argument in one native diagram.

    A node whose label names a declared entity inherits that entity's colour, so the
    code the audience learned earlier still reads on the last slide.
    """
    theme = ctx.theme
    layout = theme.layout
    nodes = ctx.spec.content["nodes"]
    flow = ctx.spec.content.get("flow", "linear")
    if flow == "hub":
        span = layout.spoke_height_in + layout.hub_height_in + theme.spacing.panel_gap_in * 2
        _summary_hub(ctx, _place(theme, ctx.body, span), nodes, ctx.role)
        return
    columns = ctx.body.columns(len(nodes), theme.spacing.panel_gap_in)
    floor = _fill_floor(ctx)
    height = max(
        _card_height(
            ctx, columns[0].width, node.get("label", ""), node.get("text", ""), floor=floor
        )
        for node in nodes
    )
    block = _place(theme, ctx.body, max(height, layout.node_min_height_in))
    for index, (node, column) in enumerate(zip(nodes, columns), start=1):
        role = ctx.deck.entities.get(node.get("label", ""), ctx.role)
        target = Box(column.left, block.top, column.width, block.height)
        add_shape(
            ctx.slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            target,
            role="card",
            detail=f"node{index}",
            fill=theme.tint(role, "panel"),
            line=theme.color(role),
            adjustment=layout.card_round_adjust,
        )
        pad = theme.spacing.card_pad_in
        add_textbox(
            ctx.slide,
            Box(
                target.left + pad, target.top + pad, target.width - 2 * pad, target.height - 2 * pad
            ),
            theme,
            [
                TextBlock(text=node.get("label", ""), step="small", bold=True, align="center"),
                TextBlock(
                    text=node.get("text", ""), step="micro", color="secondary_text", align="center"
                ),
            ],
            role="label",
            detail=f"node{index}",
            anchor="middle",
        )
        if index < len(nodes):
            mid_y = target.top + target.height / 2
            add_connector(
                ctx.slide,
                (target.right + layout.connector_gap_in, mid_y),
                (columns[index].left - layout.connector_gap_in, mid_y),
                color=theme.color("secondary_text"),
            )


def _summary_hub(ctx: SlideContext, box: Box, nodes: list[dict], default_role: str) -> None:
    theme = ctx.theme
    layout = theme.layout
    pad = theme.spacing.card_pad_in
    hub = Box(
        box.left + (box.width - layout.hub_width_in) / 2,
        box.bottom - layout.hub_height_in,
        layout.hub_width_in,
        layout.hub_height_in,
    )
    spokes = nodes[1:]
    lanes = Box(box.left, box.top, box.width, layout.spoke_height_in).columns(
        max(len(spokes), 1), theme.spacing.card_gap_in
    )
    for index, (node, lane) in enumerate(zip(spokes, lanes), start=1):
        role = ctx.deck.entities.get(node.get("label", ""), default_role)
        add_shape(
            ctx.slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            lane,
            role="card",
            detail=f"spoke{index}",
            fill=theme.tint(role, "panel"),
            line=theme.color(role),
            adjustment=layout.card_round_adjust,
        )
        add_textbox(
            ctx.slide,
            lane.inset(pad),
            theme,
            [TextBlock(text=node.get("label", ""), step="micro", bold=True, align="center")],
            role="label",
            detail=f"spoke{index}",
            anchor="middle",
        )
        add_connector(
            ctx.slide,
            (lane.left + lane.width / 2, lane.bottom),
            (hub.left + hub.width / 2, hub.top),
            color=theme.color("secondary_text"),
        )
    add_shape(
        ctx.slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        hub,
        role="card",
        detail="hub",
        fill=theme.color("primary"),
        line=None,
        adjustment=layout.card_round_adjust,
    )
    add_textbox(
        ctx.slide,
        hub.inset(pad),
        theme,
        [
            TextBlock(
                text=nodes[0].get("label", ""),
                step="small",
                bold=True,
                align="center",
                color="on_primary",
            )
        ],
        role="label",
        detail="hub",
        anchor="middle",
    )


ARCHETYPE_RENDERERS: dict[str, Callable[[SlideContext], None]] = {
    "title": archetype_title,
    "agenda": archetype_agenda,
    "section": archetype_section,
    "bullets": archetype_bullets,
    "cards": archetype_cards,
    "process": archetype_process,
    "panels": archetype_panels,
    "checklist": archetype_checklist,
    "table": archetype_table,
    "artifact": archetype_artifact,
    "statement": archetype_statement,
    "summary": archetype_summary,
}


def render_slide(ctx: SlideContext) -> None:
    """Dispatch a slide to its archetype renderer."""
    renderer = ARCHETYPE_RENDERERS.get(ctx.spec.archetype)
    if renderer is None:
        raise ValueError(f"no renderer for archetype {ctx.spec.archetype!r}")
    renderer(ctx)
