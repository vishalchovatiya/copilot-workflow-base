"""Harvest a reusable template and a design report out of an existing ``.pptx``.

A generated deck is only on-brand if it starts from a real branded package: the stock
``Presentation()`` template carries no corporate master, no corporate fonts and no
logo. This module turns any existing deck into that starting point — slides removed,
master, layouts, theme and logo kept — and writes down what it found so a theme author
can name the layouts and copy the palette without opening PowerPoint.
"""

from __future__ import annotations

import json
import shutil
import tempfile
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import asdict, dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

__all__ = [
    "HarvestError",
    "LayoutInfo",
    "MasterImage",
    "TemplateReport",
    "hydrated",
    "inspect_package",
    "make_template",
    "report_markdown",
]

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_THEME_SLOTS = (
    "dk1",
    "lt1",
    "dk2",
    "lt2",
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
    "hlink",
    "folHlink",
)


class HarvestError(Exception):
    """Raised when a source package cannot be read or stripped."""


@dataclass
class Placeholder:
    """One placeholder on a layout, with its effective geometry in inches."""

    idx: int
    ph_type: str
    name: str
    left_in: float | None
    top_in: float | None
    width_in: float | None
    height_in: float | None
    inherited: bool


@dataclass
class LayoutInfo:
    """One slide layout: how to address it and what it offers."""

    index: int
    name: str
    master: str
    used_by_slides: int
    placeholders: list[Placeholder] = field(default_factory=list)


@dataclass
class MasterImage:
    """A picture living on the master or a layout — usually a logo or wordmark."""

    owner: str
    name: str
    left_in: float
    top_in: float
    width_in: float
    height_in: float


@dataclass
class TemplateReport:
    """Everything a theme author needs to target a harvested template."""

    source: str
    slide_count: int
    layout_count: int
    width_in: float
    height_in: float
    theme_name: str
    colors: dict[str, str]
    fonts: dict[str, str]
    layouts: list[LayoutInfo]
    images: list[MasterImage]

    def to_json(self) -> str:
        """Serialise the report for machine consumption."""
        return json.dumps(asdict(self), indent=2, sort_keys=False)


@contextmanager
def hydrated(source: Path) -> Iterator[Path]:
    """Yield a guaranteed-local copy of ``source``.

    Files synced from cloud storage are frequently on-demand placeholders that raise
    ``PackageNotFoundError`` when opened in place, so every read goes through a copy in
    a temporary directory that is removed on exit.
    """
    if not source.is_file():
        raise HarvestError(f"source presentation not found: {source}")
    with tempfile.TemporaryDirectory(prefix="deckforge-") as tmp:
        local = Path(tmp) / source.name
        shutil.copyfile(source, local)
        if local.stat().st_size == 0:
            raise HarvestError(f"source presentation is empty (cloud placeholder?): {source}")
        yield local


def _theme_element(presentation) -> etree._Element | None:
    try:
        part = presentation.slide_masters[0].part.part_related_by(RT.THEME)
    except (KeyError, IndexError):
        return None
    return etree.fromstring(part.blob)


def _color_from(node: etree._Element) -> str | None:
    srgb = node.find(f"{_A}srgbClr")
    if srgb is not None:
        return str(srgb.get("val", "")).upper()
    sys_clr = node.find(f"{_A}sysClr")
    if sys_clr is not None:
        return str(sys_clr.get("lastClr", "")).upper()
    return None


def _theme_colors(theme: etree._Element | None) -> dict[str, str]:
    if theme is None:
        return {}
    scheme = theme.find(f".//{_A}clrScheme")
    if scheme is None:
        return {}
    colors: dict[str, str] = {}
    for slot in _THEME_SLOTS:
        node = scheme.find(f"{_A}{slot}")
        if node is not None:
            value = _color_from(node)
            if value:
                colors[slot] = value
    return colors


def _theme_fonts(theme: etree._Element | None) -> dict[str, str]:
    if theme is None:
        return {}
    fonts: dict[str, str] = {}
    for key, tag in (("major", "majorFont"), ("minor", "minorFont")):
        node = theme.find(f".//{_A}{tag}/{_A}latin")
        if node is not None:
            fonts[key] = str(node.get("typeface", ""))
    return fonts


def _inches(value: Emu | None) -> float | None:
    return None if value is None else round(value.inches, 3)


def _placeholders(layout) -> list[Placeholder]:
    found: list[Placeholder] = []
    for shape in layout.placeholders:
        fmt = shape.placeholder_format
        spec_pr = shape._element.spPr
        inherited = spec_pr is None or spec_pr.xfrm is None
        found.append(
            Placeholder(
                idx=fmt.idx,
                ph_type=str(fmt.type).split(" ")[0] if fmt.type is not None else "BODY",
                name=shape.name,
                left_in=_inches(shape.left),
                top_in=_inches(shape.top),
                width_in=_inches(shape.width),
                height_in=_inches(shape.height),
                inherited=inherited,
            )
        )
    return sorted(found, key=lambda p: p.idx)


def _images(container, owner: str) -> list[MasterImage]:
    found: list[MasterImage] = []
    for shape in container.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        found.append(
            MasterImage(
                owner=owner,
                name=shape.name,
                left_in=_inches(shape.left) or 0.0,
                top_in=_inches(shape.top) or 0.0,
                width_in=_inches(shape.width) or 0.0,
                height_in=_inches(shape.height) or 0.0,
            )
        )
    return found


def inspect_package(pptx_path: Path) -> TemplateReport:
    """Read a ``.pptx`` and describe its design system without modifying it."""
    with hydrated(pptx_path) as local:
        presentation = Presentation(str(local))
        theme = _theme_element(presentation)
        layout_use: dict[int, int] = {}
        layouts = list(presentation.slide_layouts)
        for slide in presentation.slides:
            if slide.slide_layout in layouts:
                layout_use[layouts.index(slide.slide_layout)] = (
                    layout_use.get(layouts.index(slide.slide_layout), 0) + 1
                )
        infos = [
            LayoutInfo(
                index=index,
                name=layout.name,
                master=presentation.slide_masters[0].name or "Master 1",
                used_by_slides=layout_use.get(index, 0),
                placeholders=_placeholders(layout),
            )
            for index, layout in enumerate(layouts)
        ]
        images = _images(presentation.slide_masters[0], "master")
        for index, layout in enumerate(layouts):
            images.extend(_images(layout, f"layout[{index}] {layout.name}"))
        return TemplateReport(
            source=pptx_path.name,
            slide_count=len(presentation.slides._sldIdLst),
            layout_count=len(infos),
            width_in=round(Emu(presentation.slide_width).inches, 3),
            height_in=round(Emu(presentation.slide_height).inches, 3),
            theme_name=(theme.get("name", "") if theme is not None else ""),
            colors=_theme_colors(theme),
            fonts=_theme_fonts(theme),
            layouts=infos,
            images=images,
        )


def make_template(source: Path, destination: Path) -> TemplateReport:
    """Write a slide-stripped copy of ``source`` to ``destination``.

    Every slide id is removed from the slide id list *and* its relationship dropped;
    dropping only one of the two leaves dangling parts behind. Media that no slide
    references any more is not written on save, so a large source deck yields a small
    template.

    Returns:
        The design report for the harvested template.
    """
    report = inspect_package(source)
    with hydrated(source) as local:
        presentation = Presentation(str(local))
        slide_ids = presentation.slides._sldIdLst
        for slide_id in list(slide_ids):
            presentation.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)
        destination.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(str(destination))
    return report


def _table(rows: list[list[str]], header: list[str]) -> list[str]:
    lines = ["| # | " + " | ".join(header) + " |"]
    lines.append("| --- | " + " | ".join("---" for _ in header) + " |")
    for number, row in enumerate(rows, start=1):
        lines.append(f"| {number} | " + " | ".join(row) + " |")
    return lines


def report_markdown(report: TemplateReport, *, template_path: Path | None = None) -> str:
    """Render the harvest report as Markdown for a human to read and act on."""
    out: list[str] = [
        f"# Template Report — {report.source}",
        "",
        "What a harvested template carries, and how to address it from a theme token",
        "file. Copy the theme colours into `palette:` and the layout indices into",
        "`layouts:` to make a new theme target this template.",
        "",
        "## Table of Contents",
        "",
        "## 1. Package",
        "",
    ]
    out += _table(
        [
            ["Source deck", report.source],
            ["Harvested template", str(template_path) if template_path else "not written"],
            ["Slides in source", str(report.slide_count)],
            ["Slides in template", "0"],
            ["Slide size (in)", f"{report.width_in} x {report.height_in}"],
            ["Theme name", report.theme_name or "unnamed"],
            ["Layouts", str(report.layout_count)],
        ],
        ["Property", "Value"],
    )
    out += ["", "## 2. Theme Colours", ""]
    if report.colors:
        out += _table(
            [[slot, f"`#{value}`"] for slot, value in report.colors.items()],
            ["Theme slot", "Hex"],
        )
    else:
        out.append("No theme part found in the package.")
    out += ["", "## 3. Fonts", ""]
    out += _table(
        [
            ["Major (headings)", report.fonts.get("major", "not set")],
            ["Minor (body)", report.fonts.get("minor", "not set")],
        ],
        ["Role", "Typeface"],
    )
    out += ["", "## 4. Layouts", ""]
    out += _table(
        [
            [str(info.index), info.name, str(len(info.placeholders)), str(info.used_by_slides)]
            for info in report.layouts
        ],
        ["Index", "Name", "Placeholders", "Used by slides"],
    )
    out += ["", "## 5. Placeholder Geometry", ""]
    geometry_rows = [
        [
            str(info.index),
            info.name,
            str(ph.idx),
            ph.ph_type,
            f"{ph.left_in}, {ph.top_in}",
            f"{ph.width_in} x {ph.height_in}",
            "inherited" if ph.inherited else "explicit",
        ]
        for info in report.layouts
        for ph in info.placeholders
    ]
    out += _table(
        geometry_rows,
        ["Layout", "Name", "idx", "Type", "Left, Top (in)", "Size (in)", "Origin"],
    )
    out += ["", "## 6. Master and Layout Images", ""]
    if report.images:
        out += _table(
            [
                [
                    image.owner,
                    image.name,
                    f"{image.left_in}, {image.top_in}",
                    f"{image.width_in} x {image.height_in}",
                ]
                for image in report.images
            ],
            ["Owner", "Shape", "Left, Top (in)", "Size (in)"],
        )
        out += [
            "",
            "Treat every master-level image as a no-draw zone: set `grid.logo_guard_in`",
            "and `grid.title_right_guard_in` so no generated shape lands on top of one.",
        ]
    else:
        out.append("No pictures on the master or layouts.")
    out.append("")
    return "\n".join(out)
